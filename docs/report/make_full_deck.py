"""Bộ slide đầy đủ 67 trang, TIẾNG VIỆT (bản để người dùng kiểm tra trước) -> docs/report/presentation_full_vi.pptx

Nội dung theo docs/report/presentation_outline.md (cùng số slide 1.1 … 6.8); lời thoại (Notes) lấy nguyên văn phần
"Speaker notes (VI)" của outline. Hình: docs/report/fig/explain/ (make_explain_figs.py), docs/report/fig/outline_vi/
(VI=1 make_outline_figs.py), docs/report/fig/paint_matching_vi.png (VI=1 make_paint_figure.py).
Chạy trên máy: python3 docs/report/make_full_deck.py
"""
import os
import re

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

D = os.path.dirname(os.path.abspath(__file__))
E = os.path.join(D, "fig", "explain")
O = os.path.join(D, "fig", "outline_vi")
PAINT = os.path.join(D, "fig", "paint_matching_vi.png")
NAVY = RGBColor(0x14, 0x24, 0x2E)
TEAL = RGBColor(0x13, 0x7C, 0x8B)
TEAL_L = RGBColor(0xD8, 0xEE, 0xF0)
AMBER = RGBColor(0xE8, 0x9B, 0x2B)
INK = RGBColor(0x1E, 0x2A, 0x32)
MUTED = RGBColor(0x5F, 0x6F, 0x7A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CARD = RGBColor(0xF3, 0xF6, 0xF8)
ICE = RGBColor(0xB9, 0xD9, 0xDE)
CODE = RGBColor(0x22, 0x2B, 0x33)
HEAD, BODY, MONO = "Cambria", "Calibri", "Consolas"
SW, SH = 13.333, 7.5
PART = {"1": "Phần 1 – Mở đầu", "2": "Phần 2 – Pipeline", "3": "Phần 3 – Chiến thuật thi đấu", "4": "Phần 4 – Vấn đề và cách sửa",
        "5": "Phần 5 – Chiến thuật từng track", "6": "Phần 6 – Kết quả và tổng kết"}

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(SW), Inches(SH)
BLANK = prs.slide_layouts[6]


# ------------------------------------------------------------------ notes from the outline
def outline_notes():
    txt = open(os.path.join(D, "presentation_outline.md"), encoding="utf-8").read()
    out = {}
    for m in re.finditer(r"^## Slide (\d\.\d+) – .*?(?=^## |\Z)", txt, re.S | re.M):
        sec = m.group(0)
        n = re.search(r"\*\*Speaker notes \(VI\):\*\*(.*?)(?=\n\n\*\*|\n## |\n---|\Z)", sec, re.S)
        if not n:
            n = re.search(r"\*\*Prepared answers \(VI[^)]*\):\*\*(.*?)(?=\n## |\n---|\Z)", sec, re.S)
        t = n.group(1) if n else ""
        t = re.sub(r"\s*\n\s*", " ", t).strip()
        t = t.replace("`", "").replace("**", "")
        out[m.group(1)] = t
    return out


NOTES = outline_notes()


# ------------------------------------------------------------------ primitives
def bg(s, c):
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = c


def tbox(s, x, y, w, h, paras, size=16, color=INK, bold=False, font=BODY, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         space=6):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, 0)
    paras = paras if isinstance(paras, list) else [paras]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space)
        segs = para if isinstance(para, list) else [para]
        for seg in segs:
            t, o = (seg, {}) if isinstance(seg, str) else seg
            r = p.add_run()
            r.text = t
            r.font.name = o.get("font", font)
            r.font.size = Pt(o.get("size", size))
            r.font.bold = o.get("bold", bold)
            r.font.color.rgb = o.get("color", color)
    return tb


def rich(line):
    """'**đậm** thường' -> runs"""
    parts = re.split(r"(\*\*[^*]+\*\*)", line)
    return [(p[2:-2], {"bold": True, "color": NAVY}) if p.startswith("**") else p for p in parts if p]


def bullets(s, x, y, w, h, items, size=None):
    n = sum(len(i) for i in items)
    if size is None:
        size = 18 if n < 260 else 16 if n < 420 else 14 if n < 620 else 12
    paras = [[("• ", {"color": TEAL, "bold": True})] + rich(i) for i in items]
    return tbox(s, x, y, w, h, paras, size=size, space=8)


def card(s, x, y, w, h, fill=CARD):
    c = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    c.adjustments[0] = 0.06
    c.fill.solid()
    c.fill.fore_color.rgb = fill
    c.line.fill.background()
    c.shadow.inherit = False
    return c


def picture(s, path, x, y, w, h):
    """fit inside the box, keep the aspect ratio, centred"""
    iw, ih = Image.open(path).size
    r = min(w / iw, h / ih)
    pw, ph = iw * r, ih * r
    s.shapes.add_picture(path, Inches(x + (w - pw) / 2), Inches(y + (h - ph) / 2), Inches(pw), Inches(ph))


def table(s, x, y, w, rows, size=12, colw=None, h_row=0.36):
    nr, nc = len(rows), len(rows[0])
    t = s.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(h_row * nr)).table
    if colw:
        for j, cw in enumerate(colw):
            t.columns[j].width = Inches(cw)
    for i, r in enumerate(rows):
        for j, v in enumerate(r):
            c = t.cell(i, j)
            c.text = ""
            p = c.text_frame.paragraphs[0]
            for seg in rich(str(v)):
                t_, o = (seg, {}) if isinstance(seg, str) else seg
                run = p.add_run()
                run.text = t_
                run.font.size = Pt(size)
                run.font.name = BODY
                run.font.bold = (i == 0) or o.get("bold", False)
                run.font.color.rgb = WHITE if i == 0 else INK
            c.fill.solid()
            c.fill.fore_color.rgb = TEAL if i == 0 else (CARD if i % 2 else WHITE)
            c.margin_left = c.margin_right = Inches(0.06)
            c.margin_top = c.margin_bottom = Inches(0.03)
    return t


def code(s, x, y, w, h, text, size=12):
    card(s, x, y, w, h, CODE)
    tbox(s, x + 0.2, y + 0.15, w - 0.4, h - 0.3, text.split("\n"), size=size, color=RGBColor(0xE6, 0xED, 0xF0), font=MONO,
         space=2)


def header(s, sid, title, sub=None):
    tbox(s, 0.55, 0.22, 10.5, 0.3, PART[sid[0]] + f" · slide {sid}", size=11, color=TEAL, bold=True)
    tbox(s, 0.55, 0.48, 12.2, 0.75, title, size=26 if len(title) < 62 else 22, bold=True, font=HEAD, color=NAVY)
    if sub:
        tbox(s, 0.55, 1.2, 12.2, 0.4, sub, size=14, color=MUTED)


N_SLIDE = [0]


def footer(s):
    N_SLIDE[0] += 1
    tbox(s, 0.55, 7.08, 8, 0.25, "SDC 2026 Midterm01 · Đội 11 · bản tiếng Việt để kiểm tra", size=9, color=MUTED)
    tbox(s, 12.2, 7.08, 0.6, 0.25, str(N_SLIDE[0]), size=9, color=MUTED, align=PP_ALIGN.RIGHT)


def new(sid, title, sub=None):
    s = prs.slides.add_slide(BLANK)
    bg(s, WHITE)
    header(s, sid, title, sub)
    footer(s)
    s.notes_slide.notes_text_frame.text = NOTES.get(sid, "")
    return s


TOP = 1.45   # content top


# ------------------------------------------------------------------ layouts
def L_fig(sid, title, items, fig, sub=None, split=4.7):
    """bullets left, figure right"""
    s = new(sid, title, sub)
    bullets(s, 0.55, TOP + (0.25 if sub else 0), split - 0.2, 5.4, items)
    picture(s, fig, split + 0.55, TOP, SW - split - 1.0, 5.5)
    return s


def L_wide(sid, title, items, fig, sub=None, text_h=1.45):
    """short bullets on top, wide figure below"""
    s = new(sid, title, sub)
    y = TOP + (0.3 if sub else 0)
    bullets(s, 0.55, y, 12.2, text_h, items, size=15 if sum(map(len, items)) > 330 else 16)
    picture(s, fig, 0.55, y + text_h + 0.05, 12.2, 7.0 - (y + text_h + 0.1))
    return s


def L_big(sid, title, fig, caption=None, sub=None):
    s = new(sid, title, sub)
    y = TOP + (0.3 if sub else 0)
    h = 7.0 - y - (0.45 if caption else 0.05)
    picture(s, fig, 0.55, y, 12.2, h)
    if caption:
        tbox(s, 0.55, 6.6, 12.2, 0.4, [rich(caption)], size=13, color=INK)
    return s


def L_two(sid, title, items, figs, sub=None, vertical=False, text_h=1.0):
    s = new(sid, title, sub)
    y = TOP + (0.3 if sub else 0)
    if items:
        bullets(s, 0.55, y, 12.2, text_h, items, size=15)
        y += text_h + 0.05
    if vertical:
        h = (7.0 - y) / 2 - 0.05
        picture(s, figs[0], 0.55, y, 12.2, h)
        picture(s, figs[1], 0.55, y + h + 0.1, 12.2, h)
    else:
        w = 6.0
        for i, f in enumerate(figs):
            picture(s, f, 0.55 + i * (w + 0.2), y, w, 7.0 - y)
    return s


def L_table(sid, title, rows, items=None, fig=None, sub=None, size=12, colw=None, fig_w=5.2):
    s = new(sid, title, sub)
    y = TOP + (0.3 if sub else 0)
    tw = 12.2 - (fig_w + 0.3 if fig else 0)
    table(s, 0.55, y, tw, rows, size=size, colw=colw)
    if items:
        bullets(s, 0.55, y + 0.38 * len(rows) + 0.3, tw, 7.0 - (y + 0.38 * len(rows) + 0.3), items, size=14)
    if fig:
        picture(s, fig, 0.55 + tw + 0.3, y, fig_w, 7.0 - y)
    return s


def L_text(sid, title, items, sub=None, code_text=None, code_h=1.6, size=None):
    s = new(sid, title, sub)
    y = TOP + (0.3 if sub else 0)
    h = 7.0 - y - ((code_h + 0.2) if code_text else 0)
    if size is None:
        n = sum(len(i) for i in items) * (1.6 if code_text else 1.0)
        size = 22 if n < 380 else 20 if n < 560 else 18 if n < 760 else 16
    bullets(s, 0.55, y, 12.2, h, items, size=size)
    if code_text:
        code(s, 0.55, 7.0 - code_h, 12.2, code_h - 0.05, code_text)
    return s


def problem(sid, title, symptom, cause, fix, result, fig=None, figs=None, vertical=False, code_text=None):
    """Phần 4: triệu chứng → nguyên nhân → cách sửa → trước / sau"""
    s = new(sid, title)
    parts = (("Triệu chứng", symptom, AMBER), ("Nguyên nhân", cause, MUTED), ("Cách sửa", fix, TEAL), ("Trước → sau", result, NAVY))
    wide = False
    if fig:
        iw, ih = Image.open(fig).size
        wide = iw / ih >= 1.7
    if wide:
        # 2 × 2 cards on top, the figure below over the full width
        longest = max(len(t) for _, t, _ in parts)
        hh = 0.62 + 0.24 * (longest // 85)
        for i, (lab, txt, col) in enumerate(parts):
            x, y = 0.55 + (i % 2) * 6.15, TOP + (i // 2) * (hh + 0.08)
            card(s, x, y, 6.05, hh)
            tbox(s, x + 0.15, y + 0.05, 5.8, 0.3, lab, size=11, bold=True, color=col)
            tbox(s, x + 0.15, y + 0.29, 5.8, hh - 0.3, [rich(txt)], size=12.5)
        y = TOP + 2 * (hh + 0.08)
        picture(s, fig, 0.55, y, 12.2, 7.0 - y)
        return s
    has_fig = fig or figs
    w = 4.6 if has_fig else 12.2
    y = TOP
    for lab, txt, col in parts:
        hh = 0.55 + 0.25 * (len(txt) // (52 if has_fig else 140))
        card(s, 0.55, y, w, hh)
        tbox(s, 0.7, y + 0.06, w - 0.3, 0.3, lab, size=11, bold=True, color=col)
        tbox(s, 0.7, y + 0.3, w - 0.3, hh - 0.3, [rich(txt)], size=12.5 if has_fig else 15)
        y += hh + 0.1
    if code_text:
        code(s, 0.55, y, w, min(1.4, 7.0 - y), code_text, size=11)
    if fig:
        picture(s, fig, 5.4, TOP, 7.4, 5.55)
    if figs:
        if vertical:
            picture(s, figs[0], 5.4, TOP, 7.4, 2.72)
            picture(s, figs[1], 5.4, TOP + 2.83, 7.4, 2.72)
        else:
            picture(s, figs[0], 5.4, TOP, 3.65, 5.5)
            picture(s, figs[1], 9.15, TOP, 3.65, 5.5)
    return s


def boxes_diagram(s, rows, y0, label_w=1.9):
    """rows: [(label, [box texts], color)] -> chain of boxes with arrows"""
    for r, (lab, items, colr) in enumerate(rows):
        y = y0 + r * 2.55
        card(s, 0.55, y, 12.2, 2.3, TEAL_L if r == 0 else CARD)
        tbox(s, 0.75, y + 0.15, label_w, 2.0, lab, size=15, bold=True, color=NAVY, anchor=MSO_ANCHOR.MIDDLE)
        n = len(items)
        bw = (12.2 - label_w - 0.6 - 0.35 * (n - 1)) / n
        for i, t in enumerate(items):
            x = 0.55 + label_w + 0.4 + i * (bw + 0.35)
            b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y + 0.35), Inches(bw), Inches(1.6))
            b.fill.solid()
            b.fill.fore_color.rgb = colr
            b.line.fill.background()
            tf = b.text_frame
            tf.word_wrap = True
            tf.margin_left = tf.margin_right = Inches(0.06)
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            run = p.add_run()
            run.text = t
            run.font.size, run.font.name = Pt(11.5), BODY
            run.font.color.rgb = WHITE
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            if i:
                c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x - 0.33), Inches(y + 1.15), Inches(x - 0.02),
                                           Inches(y + 1.15))
                c.line.color.rgb = NAVY
                c.line.width = Pt(2)
                c.line.end_arrowhead = True if hasattr(c.line, "end_arrowhead") else None


def dark(sid, title, lines, sub=None, fig=None):
    s = prs.slides.add_slide(BLANK)
    bg(s, NAVY)
    tbox(s, 0.8, 1.4 if fig else 2.0, 7.4 if fig else 11.6, 1.6, title, size=36, bold=True, font=HEAD, color=WHITE)
    if sub:
        tbox(s, 0.8, 3.15, 7.4 if fig else 11.6, 1.0, sub, size=18, color=ICE)
    tbox(s, 0.8, 4.6, 7.4 if fig else 11.6, 2.2, lines, size=15, color=AMBER, bold=True)
    if fig:
        picture(s, fig, 8.5, 0.9, 4.4, 5.8)
    footer(s)
    s.notes_slide.notes_text_frame.text = NOTES.get(sid, "")
    return s


# =================================================================== PHẦN 1
dark("1.1", "Định vị xe trên bản đồ LiDAR",
     ["SDC 2026 · Midterm01 · Đội 11 · [CẦN ĐIỀN: họ tên] · MSSV 114550202",
      "Làm một mình, có trợ lý lập trình AI (Claude Code)",
      "Bảng điểm public 25/10: Track01 hạng 1 (0.36) · Track02 hạng 1 (0.11) · Track03 hạng 2 (0.28)"],
     sub="Point cloud dựng từ ảnh, vạch sơn mặt đường và IMU – Track01, Track02 và Track03 không có GNSS",
     fig=os.path.join(O, "trajectory_on_map_t3.png"))

L_wide("1.2", "Bài toán: mỗi frame camera một pose (x, y, yaw)",
       ["**Đầu vào:** point cloud dựng từ ảnh mỗi frame (raw_ply / 3cameras_ply), bản đồ LiDAR map.pcd (x y z cường độ), "
        "init_guess.txt (pose đầu có nhiễu), rosbag: IMU 100 Hz, 3 luồng video h265, GNSS (chỉ Track01/02)",
        "**Đầu ra:** CSV idx, x, y, yaw (radian). **Điểm:** E = RMSE(Δp) + 0.2·RMSE(Δθ) → sai 1° yaw = sai 0.2 m vị trí",
        "**Luật 10/22 cho Track03:** chỉ dùng frame hiện tại và quá khứ, không làm mượt ngược, TA chạy lại code"],
       os.path.join(E, "x01_t1_cameras.png"), text_h=1.9)

L_wide("1.3", "Hai loại point cloud từ ảnh – không cái nào đúng đơn vị mét",
       ["**raw_ply (LoGeR, 1 camera):** 20–41 nghìn điểm/frame, mặt đường chỉ 0.18 đơn vị dưới camera → phải nhân ≈ 9",
        "**3cameras_ply (MapAnything, 3 camera):** ~330 nghìn điểm, gần đúng mét (×0.8) – **Track03 không có loại này**",
        "Cả hai đều méo: tỉ lệ dao động 5–7 % giữa các frame, độ sâu càng xa càng sai"],
       os.path.join(E, "x02_t1_clouds.png"), text_h=1.5)

s = L_table("1.4", "Ba track, ba kiểu khó",
            [["", "Track01", "Track02", "Track03"],
             ["Số frame", "3011", "3299", "3142"],
             ["GNSS", "có, giật 0.142 m", "có, mượt 0.052 m", "**không có**"],
             ["Chấm theo", "public toàn track", "public toàn track", "**private ≈ 87 %**"],
             ["Đặc biệt", "nhà / cây che 74 % đường, 11 bước nhảy GNSS", "đứng yên 13 s đầu, init lệch 14.85°",
              "đèn đỏ 43 s = 44 % private, rẽ trái 88°"],
             ["Baseline TA", "1.79075", "1.92587", "private ≤ 1.40"]],
            size=12, colw=[1.6, 3.5, 3.5, 3.6])
picture(s, os.path.join(O, "t3_segments.png"), 0.55, 4.0, 12.2, 2.95)

L_fig("1.5", "Những gì làm bài này khó",
      ["**Sai tỉ lệ trong code mẫu** (×3, đúng ≈ 9) → NDT trôi 91 m / 110 m",
       "**Point cloud AI bị méo:** trên point cloud hoàn hảo GICP chỉ sai 5–8 mm; trên point cloud thật sai ~0.8 m mỗi frame → "
       "lỗi nằm ở dữ liệu, không ở thuật toán khớp (hình)",
       "**Khởi đầu có nhiễu:** init lệch 2° (T1), 14.85° (T2), roll/pitch ~15° (T3)",
       "**Yaw GNSS** tính từ vị trí: trễ 1.1–1.4 s, đứng im khi xe dừng",
       "**Track03 không GNSS:** odometry từ point cloud chỉ bằng 0.12–1.16 lần quãng đường thật",
       "Vật thể động: không đo riêng [CẦN ĐIỀN nếu muốn nêu]"],
      os.path.join(E, "x16_synthetic_gicp.png"))

# =================================================================== PHẦN 2
s = new("2.1", "Toàn bộ pipeline")
boxes_diagram(s, [("Track01 / 02\n(có GNSS)", ["GNSS 4 Hz", "Kalman CTRV + RTS\n(nhiễu theo maximum likelihood)",
                                               "3cameras_ply: tỉ lệ = 1.646 m / độ cao mặt đường, tầm ≤ 20 m",
                                               "GICP với bản đồ 0.3 m\n2 lượt, trừ lệch TB", "CSM bản đồ + yaw = gyro + bản đồ",
                                               "result_1 / result_2"], TEAL),
                  ("Track03\n(không GNSS,\nnhân quả)", ["raw_ply + IMU", "Bộ bám: MHT, lọc tốc độ IMU, gyro theo trọng lực, ZUPT",
                                                        "CSM ±15 m rồi GICP", "Ảnh h265: chiếu mặt đường (IPM), tách vạch sơn",
                                                        "NCC với vạch sơn trong cường độ LiDAR", "Bộ lọc + chuyển động IMU → t3_Q4"],
                   RGBColor(0x2A, 0x78, 0xD6))], TOP + 0.1)

s = new("2.2", "Code mẫu so với phần nhóm tự làm")
card(s, 0.55, TOP, 5.9, 5.45)
tbox(s, 0.8, TOP + 0.15, 5.4, 0.4, "Code mẫu (localization.cpp, map_modified.cpp)", size=16, bold=True, color=MUTED)
bullets(s, 0.8, TOP + 0.7, 5.4, 4.6, ["NDT của PCL, ply_scale 3", "voxel đầu vào 0.3 m, bản đồ 1 m, vùng bản đồ ±100 m",
                                      "NDT: ô 8 m, bước 0.8, ε 0.05, 30 vòng lặp",
                                      "frame sau bắt đầu từ pose frame trước; frame đầu từ init_guess",
                                      "kết quả chạy mặc định: trôi, nộp được 602.39 điểm"], size=15)
card(s, 6.85, TOP, 5.9, 5.45, TEAL_L)
tbox(s, 7.1, TOP + 0.15, 5.4, 0.4, "Nhóm tự làm", size=16, bold=True, color=TEAL)
bullets(s, 7.1, TOP + 0.7, 5.4, 4.6, ["**C++** (~130 dòng): prior từng frame, loại kết quả nhảy xa, chạy cách frame, 10 đoạn song song, file debug",
                                       "**59 công cụ Python** (~6 600 dòng): tỉ lệ theo mặt đường, Kalman/RTS, GICP, lọc yaw, CSM + yaw bản đồ",
                                       "**Bộ bám không GNSS** cho Track03 (MHT, lọc tốc độ IMU, ZUPT, CSM)",
                                       "**Khớp vạch sơn từ camera** (giải mã h265, IPM, NCC qua FFT, bộ lọc + IMU) – đưa Track03 từ 0.80 xuống 0.28",
                                       "Công cụ kiểm tra không cần đáp án, bot Kaggle, mọi file nộp tạo lại được từng byte"], size=14)

L_wide("2.3", "Bản đồ: cỡ voxel, vùng cắt và các bản đồ phụ",
       ["Code mẫu giảm mẫu 1 m (59.5 triệu → 0.5 triệu điểm); pipeline của nhóm dùng **0.3 m** cho GICP – thử 0.2 / 0.3 / 0.5 m: nhiễu chỉ đổi ≤ 0.04 m, 0.2 m chậm gấp 3",
        "Thứ quan trọng hơn là **tầm point cloud ≤ 20 m**. Bản đồ phụ: trường xác suất (vật cao 0.4–3 m) cho CSM; raster vạch sơn (cường độ LiDAR, white top-hat 0.7 m)"],
       os.path.join(E, "x13_map_voxel.png"), text_h=1.35)

L_wide("2.4", "Xử lý point cloud từ ảnh: tỉ lệ, tầm, bộ lọc",
       ["**Tỉ lệ mỗi frame** = 1.646 m / độ cao mặt đường (trung vị 57 frame ~2 s; Track03 chỉ dùng frame đã qua). **Tầm ≤ 20 m:** Track01 0.80 → 0.63",
        "Mặt nạ capo: point cloud 3 camera chỉ 0.5–2 % điểm trong vùng capo → không cần cho GICP, **có dùng** khi chiếu ảnh. "
        "Lọc điểm lạc: kém đi (0.714 → 0.764 m). Bỏ mặt đường: GICP chỉ hội tụ 30 %"],
       os.path.join(E, "x15_scale_x3.png"), text_h=1.45)

L_wide("2.5", "Khớp quét: NDT → GICP, thô rồi tinh",
       ["**NDT** sau khi sửa tỉ lệ và ô 4 m: ổn định nhưng kém GNSS. **GICP** (small_gicp): nhanh hơn ~200 lần, chính xác hơn (0.803 so với 0.836 m); "
        "chỉ **2 lượt** và **trừ lệch trung bình**",
        "Hình: cùng một frame Track01, màu = khoảng cách từ điểm point cloud ảnh tới bản đồ. Track03 khớp thô bằng **CSM ±15 m** rồi mới GICP; "
        "tầng cuối là **khớp vạch sơn**"],
       os.path.join(E, "x03_t1_gicp.png"), text_h=1.45)

L_two("2.6", "Cập nhật trạng thái: bộ lọc, IMU, khi khớp thất bại",
      ["**Track01/02:** EKF + RTS (CTRV, nhiễu theo maximum likelihood) – **có làm mượt ngược, không nhân quả**. "
       "**Track03 (nhân quả):** Kalman vị trí, cổng 3σ; lọc tốc độ IMU; ZUPT; gyro chiếu theo trọng lực",
       "Khớp bị loại → chỉ dự đoán; đứng yên → không cập nhật GICP; khởi đầu: nhiều giả thuyết yaw, mỗi frame xuất giả thuyết tốt nhất tại frame đó"],
      [os.path.join(E, "x12_t3_mht.png"), os.path.join(E, "x09_t3_zupt.png")], text_h=1.45)

# =================================================================== PHẦN 3
s = new("3.1", "Cách làm: từ baseline, mỗi lần đổi một thứ")
steps = ["Giả thuyết\nghi trước", "Đổi đúng\n1 tham số", "Kiểm tra offline\n(bản đồ, IMU, Track01/02)", "Nộp Kaggle\n(≤ 3 file, được duyệt)",
         "Ghi nhật ký\n+ kết luận"]
for i, t in enumerate(steps):
    x = 0.7 + i * 2.5
    b = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(TOP + 0.4), Inches(2.0), Inches(2.0))
    b.fill.solid()
    b.fill.fore_color.rgb = TEAL if i % 2 == 0 else NAVY
    b.line.fill.background()
    tf = b.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = t
    r.font.size, r.font.color.rgb, r.font.bold = Pt(13), WHITE, True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    if i < 4:
        tbox(s, x + 2.02, TOP + 1.15, 0.45, 0.5, "→", size=28, color=AMBER, bold=True)
bullets(s, 0.6, TOP + 2.8, 12.1, 2.8, ["Ngày đầu: chạy code mẫu → trôi, nộp được **602.39**; nộp thử GNSS → **1.31 / 0.57** → GNSS làm nền cho Track01/02",
                                         "Không có đáp án → thước đo offline: độ khớp bản đồ (Spearman với Kaggle +0.92 / +0.67), tích phân IMU, "
                                         "chạy thuật toán Track03 trên Track01/02 như không có GNSS, so camera trước với camera bên",
                                         "Bảng điểm chỉ dùng như phép đo ở Track01/02 (TA cho phép) – không bao giờ ở Track03"], size=15)

s = L_text("3.2", "Công cụ: bot Kaggle, Claude Code, quản lý phiên bản",
           ["**Bot Kaggle** (mochihat/kaggle-sdc-bot, chỉ thư viện chuẩn): status, check (số dòng, yaw radian, khoảng cách tới quỹ đạo tham chiếu → "
            "bắt nộp nhầm track), submit, run (cron 30 phút, ≤ 3 file/lần, 15/ngày)",
            "**Claude Code:** brief đánh số nhiệm vụ (T0–T13), CLAUDE.md để bàn giao giữa các phiên, nhật ký 13 file, commit sau mỗi nhiệm vụ; "
            "AI không bao giờ nộp khi chưa có \"cho phép nộp\"",
            "**Git:** 97+ commit, commit đầu = code mẫu nguyên gốc; mọi file đã nộp được lưu và tạo lại được từng byte"],
           code_text="$ python3 bot.py status\n== Track3 (sdc-2026-localization-track-03)\n   1. team 12   0.24000\n   2. team 11   0.28000   <== us\n"
                     "   3. team 3    0.49000\n$ python3 bot.py check 2 gnss_1.csv      →  (False, '3011 rows, expected 3299')",
           code_h=2.3)

L_big("3.3", "Track01/02: điểm tốt nhất theo thời gian, theo nguồn gốc", os.path.join(O, "timeline_t1t2.png"),
      caption="Pipeline thuần: **0.51 / 0.24**. Có dò bảng điểm: **0.36 / 0.11** (hạng 1 cả hai). Bậc ở 22/10 = dựng lại pipeline sạch sau khi TA trả lời.")
L_big("3.4", "Track03: 58 điểm public, chỉ output của pipeline (không dò)", os.path.join(O, "timeline_t3.png"),
      caption="16 lần đầu (cam) từ bộ bám bản đầu có phần khởi đầu chưa nhân quả – tự phát hiện và sửa ngày 23/10. Đường xanh chỉ tính các bản nhân quả.")

# =================================================================== PHẦN 4
problem("4.1", "Môi trường: code mẫu không build được",
        "Docker build lỗi, container không chạy.",
        "pandas<2.1 không đặt trong nháy (shell hiểu < là chuyển hướng); đường dẫn mẫu /{USER-PATH}/; khối GPU khi máy chưa có nvidia toolkit.",
        "Đặt \"pandas<2.1\" trong nháy, sửa đường dẫn, tắt khối GPU (NDT chạy CPU); lưu bản sửa trong docker/.",
        "build lỗi → build xong trong ~66 s, colcon build OK.",
        code_text='- RUN pip install pandas<2.1\n+ RUN pip install "pandas<2.1"\n- /{USER-PATH}/SDC-midterm01\n+ /home/thuhang/SDC-midterm01')
problem("4.2", "Tải dữ liệu bị chặn",
        "\"download all\" báo 404; tải từng file qua API dừng ở ~3 GB với lỗi 429 Too Many Requests.",
        "Cuộc thi quá lớn (Track01: 15 066 file, 25.4 GB), API giới hạn tốc độ.",
        "Tải zip trên trình duyệt, so số file / dung lượng / CRC của 13 205 file; Track03 lấy từ link Google Drive trong PDF luật.",
        "mất thời gian, không ảnh hưởng điểm.")
problem("4.3", "Code mẫu trôi 91–110 m vì sai tỉ lệ point cloud",
        "Lệch > 5 m từ frame 17 (T1) / frame 0 (T2); tới 91.2 m và 110.0 m; cả quỹ đạo nộp được 602.39.",
        "ply_scale = 3, trong khi mặt đường chỉ 0.18 đơn vị dưới camera (camera cao ~1.6 m) → tỉ lệ ≈ 9. Ở ×3 cả cảnh co lại, NDT bám vào cấu trúc gần nhất và cộng dồn.",
        "Tỉ lệ ×9, sau đó tỉ lệ từng frame = 1.646 m / độ cao mặt đường.",
        "độ sửa NDT: 4.32 m, yaw ±47° (×3) → 0.70 m, yaw ±2° (×9).", fig=os.path.join(E, "x15_scale_x3.png"))
problem("4.4", "Lần kiểm tra tỉ lệ đầu tiên chỉ ra tỉ lệ sai",
        "Thước đo khoảng cách trung bình point cloud – bản đồ nhỏ nhất ở tỉ lệ 1 (lần đầu) / 3.",
        "Thước đo thiên vị: point cloud càng nhỏ càng dồn sát mặt đường nên luôn \"gần\" bản đồ.",
        "Xem độ lệch yaw tốt nhất có nhất quán không: ở ×3, 4/8 frame phải quay 24–30°; ở ×9, 7/8 frame ≤ 6°.",
        "kết luận sai (×1) → đúng (×9, trục FLU), NDT xác nhận.", fig=os.path.join(E, "x02_t1_clouds.png"))
problem("4.5", "Ô NDT quá thô",
        "Ở ×9 với ô 8 m: độ sửa trung vị 1.55 m, chỉ 35 % ≤ 1 m, yaw ±9.3° (T2).",
        "Point cloud ×9 chỉ phủ ~27 m; ô 8 m quá thô.",
        "ndt_resolution 4.",
        "1.55 m → 0.23 m, 92 % ≤ 1 m, yaw ±2.1°; ×10 kém hơn ×9.")
problem("4.6", "NDT ổn định vẫn làm điểm xấu đi",
        "Cộng độ sửa NDT vào GNSS: T2 0.47281 → 0.53520, T1 1.25338 → 1.27855.",
        "0.2–0.5 m mà NDT \"sửa\" chính là sai số của NDT trên point cloud LoGeR.",
        "Không ghép NDT raw_ply; thử point cloud 3 camera ở ×0.8: lần đầu sửa được lỗi vị trí GNSS.",
        "1.21861 → 1.15042 (Track01, chỉ đổi vị trí).", fig=os.path.join(O, "t1_position_ablation.png"))
problem("4.7", "NDT chỉ dùng ảnh trôi dần",
        "Từ init_guess, 150 frame: T1 sai 1.5 → 9.1 m; T2 3.6 → 8.9 m, yaw lệch +29…+40°.",
        "Không có mô hình chuyển động; hướng ban đầu Track02 lệch ~15°.",
        "Thành thiết kế của Track03: dự đoán bằng odometry + IMU, nhiều giả thuyết yaw lúc đầu.",
        "trôi 9 m / 150 frame → bộ bám Track03 nhận 96–99 % frame.", fig=os.path.join(E, "x12_t3_mht.png"))
problem("4.8", "Nộp nhầm track",
        "\"Submission has 3011 rows but 3299 were expected\" – mất 1 lượt.",
        "Nộp tay, không kiểm tra.",
        "bot.py check: số dòng, header, idx, giá trị hữu hạn, yaw radian, khoảng cách tới quỹ đạo tham chiếu ≤ 20 m.",
        "1 lượt mất → không còn nộp nhầm.", code_text="$ python3 bot.py check 2 gnss_1.csv\n(False, '3011 rows, expected 3299')")
problem("4.9", "Yaw GNSS trễ và đứng im",
        "Gyro khớp với yaw GNSS khi lệch −1.45 s; 13 s đầu Track02 yaw GNSS giữ −15.206°.",
        "Yaw GNSS = hướng dịch chuyển 1 s của vị trí GNSS, cập nhật ~1.25 s/lần kiểu bậc thang, đứng im khi xe dừng.",
        "Dựng lại yaw lúc đứng yên bằng gyro (hướng lúc xe đạt 4 m/s, tích phân ngược, trừ bias): −19.78° thay vì −15.21°.",
        "T2 0.34530 → 0.33504.", fig=os.path.join(E, "x10_t2_gnss_yaw.png"))
problem("4.10", "Bù độ trễ lại làm xấu đi: yaw đáp án là yaw thân xe",
        "Dùng hướng chạy đúng giờ: T2 0.40778 → 0.45496, T1 1.25338 → 1.30875.",
        "Đáp án dùng yaw thân xe, không phải hướng chuyển động; hướng chạy từ GNSS giật thì nhiễu.",
        "Bộ lọc bù: gyro trục z cho phần nhanh + làm mượt (hướng chạy − yaw gyro) với hằng số τ cho phần chậm.",
        "T1 0.92775 → 0.78914 (yaw 1.559° → 0.865°), T2 0.29887 → 0.27654. τ, trục gyro chọn theo điểm.",
        fig=os.path.join(O, "yaw_tau.png"))
problem("4.11", "Lặp GICP nhiều lượt lại xấu đi",
        "Lượt 2: 0.679 m; lượt 5: 0.801 m (lỗi vị trí Track01).",
        "Point cloud méo → cực tiểu của GICP lệch khỏi vị trí thật: mặt đường khớp (~0.2 m), nhà / cây không (0.4–1.3 m).",
        "Dừng ở 2 lượt (chọn sau khi thấy lượt 5 – tham số theo điểm).",
        "1 lượt 1.02813 → 2 lượt 0.99283; 5 lượt 1.11461.", fig=os.path.join(O, "t1_position_ablation.png"))
problem("4.12", "GICP lệch cố định",
        "Độ sửa trung bình +0.17 m dọc, −0.14 m ngang; dò cho thấy GNSS không lệch ngang (+0.025 m), không lệch giờ (−3.6 ms).",
        "Lệch của phép khớp ảnh (gốc camera / tỉ lệ của point cloud MapAnything).",
        "--demean: GNSS giữ phần trung bình, GICP chỉ thêm phần dao động.",
        "0.99060 → 0.92775 (lỗi vị trí ~0.679 → ~0.616 m, dự đoán 0.62).", fig=os.path.join(E, "x03_t1_gicp.png"))
problem("4.13", "Yaw của GICP và khớp quét ở Track02 không dùng được",
        "Thêm yaw GICP: yaw 1.570° → ~1.627°. Track02: GNSS 0.2267 m → ICP 0.444, GICP 0.414, NDT 0.270 m.",
        "Yaw từng frame của point cloud ảnh sai 2.7–2.9°; GNSS Track02 (~0.2 m) chính xác hơn phép khớp (~1 m có hệ thống).",
        "Không lấy yaw từ GICP; Track02 không dùng khớp quét cho vị trí (chỉ CSM bản đồ trọng số thấp).",
        "Track02 có GICP 0.30 so với không 0.28.", fig=os.path.join(E, "x16_synthetic_gicp.png"))
problem("4.14", "Điểm ở xa mang phần lớn sai số",
        "Tầm 40 m: nhiễu dọc / ngang 0.665 / 0.532 m, chỉ 53 % frame hợp lệ.",
        "Độ sâu từ ảnh sai tăng theo khoảng cách (thí nghiệm tổng hợp: 0.35–0.43 m).",
        "GICP --max-range 20.",
        "offline: 0.563 / 0.403 m, 58 % hợp lệ, nhanh 2.5×; Kaggle T1 0.80 → 0.63 (15 m: 0.63, 10 m: 0.67).",
        fig=os.path.join(E, "x14_t4_sweep.png"))
problem("4.15", "TA trả lời: kết quả dò không phải bài chấm",
        "File tốt nhất là tổ hợp / dò bảng điểm (T1 0.46, T2 0.19).",
        "TA (20/10): nộp nhiễu loạn có hệ thống để suy ra lỗi ẩn = leaderboard fitting, không làm cơ sở bài chấm.",
        "Dừng dò; phân loại lại mọi file (pipeline / lai / dò); dựng lại pipeline sạch bằng một lệnh, trùng từng byte.",
        "file sạch 0.80 / 0.28 (22/10). (23/10 TA cho dò ở Track01/02 – slide 5.2.)", fig=os.path.join(O, "timeline_t1t2.png"))
problem("4.16", "File \"lai\": phụ thuộc ngầm vào phần chọn theo điểm",
        "cfz1_tau7 (0.78914) khởi tạo GICP từ yaw tổ hợp; z / roll / pitch lấy từ lượt NDT cũ chạy trên prior tổ hợp.",
        "File dựng chồng lên file khác.",
        "Dựng lại sạch; z / roll / pitch từ mặt đường bản đồ + 3 hằng số hiệu chuẩn (độ cao camera, pitch lắp +2.47° / +2.30°, roll).",
        "vị trí đổi 0.060 m RMS; điểm 0.80 → 0.81 (ngang).", fig=os.path.join(O, "timeline_t1t2.png"))
problem("4.17", "GICP Track02 đọc nhầm độ cao của Track01",
        "Lượt GICP Track02 đầu tiên hỏng.",
        "--pose-source c3s10 (Track01) thay vì c3s10T2 → độ cao khởi đầu 109.9 m trên bản đồ ~17.4 m.",
        "Kiểm tra đầu vào trong gicp_refine.py: số dòng prior, nguồn pose, độ cao khởi đầu lệch > 20 m thì dừng.",
        "chạy sai mà không biết → báo lỗi ngay.",
        code_text="start height z = 109.9 m but the map around\nthe first pose is at ~17.4 m")
problem("4.18", "Bộ ghép \"thông minh\" hơn không giúp: lỗi là có hệ thống",
        "Bình phương tối thiểu theo Hessian GICP: 0.84 / 0.87 so với 0.80 (T1); T2 0.30 so với 0.28.",
        "Hessian đúng xu hướng trung bình nhưng không dự báo nhiễu từng frame (Spearman ≈ 0); nhiễu đo được đẳng hướng 0.505 / 0.494 m; phần còn lại là lệch chậm có hệ thống.",
        "Giữ bộ ghép đơn giản; sửa đầu vào (tầm 20 m). Kalman bền vững (Huber) chỉ đổi 1.9 cm → không dùng.",
        "giữ 0.80, rồi tầm 20 m → 0.63.", fig=os.path.join(E, "x16_synthetic_gicp.png"))
problem("4.19", "Tổ hợp dò bảng điểm bị suy biến",
        "Trọng số −9.2…+7.5, có lúc 10¹⁴; lb_t1_r3 dự đoán 0.33, thật 0.42.",
        "Các hướng gần trùng nhau + điểm bị cắt còn 2 chữ số (±0.005).",
        "Chính quy hoá ridge, bỏ nghiệm |a| > 20, chọn λ bằng Monte Carlo có nhiễu cắt số.",
        "0.42 → 0.40 → 0.36 (hạng 1). Đây là kỹ thuật dò, trọng số không có ý nghĩa vật lý.", fig=os.path.join(E, "x11_probe_parabola.png"))
problem("4.20", "Track03: bộ bám đầu tiên không nhân quả",
        "Bản N được 0.77 nhưng 80 frame đầu \"đẹp quá\".",
        "Khởi đầu nhiều giả thuyết: chọn giả thuyết thắng ở frame ~60–80 rồi ghi lại quỹ đạo quá khứ của nó → dùng thông tin tương lai.",
        "Mỗi frame chỉ xuất giả thuyết tốt nhất tại chính frame đó.",
        "0.77 (không hợp lệ) → 0.96 (A2h, nhân quả). Phải chọn tay file nhân quả làm bài chấm.", fig=os.path.join(E, "x12_t3_mht.png"))
problem("4.21", "Track03: mỗi khúc cua mất 14 % góc quay",
        "Khúc cua 90° chỉ thành 77°; sau cua yaw lệch −12°, quỹ đạo chạy cách đường 20 m.",
        "Tích phân gyro dừng ở mẫu IMU cuối trước frame (IMU 100 Hz, camera ~28 Hz), bỏ mất đoạn tới thời điểm frame → 86 %.",
        "Tích phân tới đúng thời điểm frame.",
        "nằm trong A2h; ảnh hưởng chủ yếu ở phần private.", fig=os.path.join(O, "t3_drift_N_vs_A2h.png"))
problem("4.22", "Track03: trôi 29 m dọc đường trước đèn đỏ",
        "Đồ thị độ khớp dọc đường: −1 m (frame 300) → −13.5 → −21 → −29 m; tích phân gia tốc kế: 75 ± 8 m, bộ bám 96 m.",
        "Odometry từ point cloud LoGeR thất thường (0.12–1.16 × quãng đường thật); GICP chỉ nhìn 1–2 m.",
        "ZUPT + bias gyro online + **khớp lưới CSM** trên trường xác suất (dọc ±15 m × ngang ±2 m), rồi GICP.",
        "GICP được chấp nhận 69 % → 98 %; quãng đường 600 → dừng: 72 m (IMU 75 ± 8).", fig=os.path.join(E, "x04_t3_csm.png"))
problem("4.23", "Track03: GICP trượt dọc đường",
        "Đổi một tham số nhỏ là lạc ở frame ~470–1000 (lệch 22–30 m, chỉ 40 % chấp nhận).",
        "Mỗi frame GICP hội tụ trước dự đoán ~1.1 m mà vẫn qua cổng → 11 m sau 12 frame.",
        "Nhiễu đo bất đẳng hướng: GICP chỉ được tin 0.7 m theo ngang, 2 m theo dọc (CSM lo phương dọc).",
        "mọi biến thể bám 96–99 % và khớp nhau ±0.5 m.", fig=os.path.join(O, "t3_drift_N_vs_A2h.png"))
problem("4.24", "Track03: tay đòn camera sửa quá tay",
        "Tay đòn CSM −1.0 m (hiệu chuẩn trên Track01/02): public 0.96 → 1.05.",
        "Point cloud LoGeR dài hơn thật 10 % ở Track01/02 nhưng chỉ 5 % ở Track03 → −1.0 m chủ yếu là hệ quả của sai tỉ lệ.",
        "Đặt tay đòn theo sai tỉ lệ đo trên chính Track03: ≈ −0.7 m.",
        "1.05 → 0.84 (L2b).")
problem("4.25", "Track03: IMU lắp nghiêng 6°",
        "Yaw sai qua khúc cua; ước lượng trọng lực nhân quả lệch 8° ở giây đầu (W3h 0.89 so với L2b 0.84).",
        "Trục z của gyro chỉ thấy cos 6° của góc quay quanh phương thẳng đứng và trộn bias trục x/y; xe vừa tăng tốc vừa rẽ ngay từ frame 0.",
        "Tốc độ quay = gyro · hướng trọng lực cố định (hiệu chuẩn lúc đứng yên) + τ yaw 30 s.",
        "0.89 → 0.80 (X3f); kiểm tra qua khúc cua 88°: lệch bản đồ +1.43° → −0.02°.")
problem("4.26", "Track03: GICP kéo vị trí khi xe đứng yên",
        "Track02: 230 frame đứng yên kéo lệch −1.0 m dọc / +1.9 m ngang; Track03: lùi 0.3–0.5 m, yaw trôi tới −2.4°.",
        "Các frame của cùng một khung cảnh có chung một sai số – cập nhật hàng trăm lần một phép đo sai.",
        "Không cập nhật vị trí / yaw bằng GICP khi ZUPT báo đứng yên.",
        "yaw giữ −0.9° suốt 45 s dừng; Track02 lúc đứng 1.95 → 1.56 m.", fig=os.path.join(O, "t3_segments.png"))
problem("4.27", "Track03: chậm 4 m khi xe chạy lại (hai lỗi của bộ lọc tốc độ)",
        "Sau 43 s đèn đỏ, vị trí dọc chậm 3–4 m tới khi CSM bắt kịp ở frame 2306–2318.",
        "(1) Số đọc odometry hỏng trông \"chính xác\" (nhiễu tỉ lệ với số đọc) kéo tốc độ về 0; (2) ZUPT tưởng tăng tốc đều (~1.2 m/s²) là đứng yên (hình trên).",
        "--speed-odo-low 0.5 và --zupt-ax 0.3 (đứng yên còn cần gia tốc tiến TB nhỏ).",
        "sai số dọc 2.46 → 0.72 m; ước tính private −0.13…−0.21; public giữ 0.80.",
        figs=[os.path.join(E, "x09_t3_zupt.png"), os.path.join(O, "t3_restart_fix.png")], vertical=True)
problem("4.28", "Track03: point cloud LoGeR chỉ cho độ chính xác ~1 m → dùng vạch sơn",
        "Dù chỉnh nhiễu thế nào, sai số khi bám vẫn 0.90–1.06 m; public kẹt ở 0.80.",
        "Độ sâu LoGeR sai theo khoảng cách; hiệu chuẩn chỉ tới ±0.3 m; Track03 không có point cloud 3 camera.",
        "Mốc mới chính xác cỡ cm: **vạch sơn** trong cường độ LiDAR, camera nhìn thấy (video h265 trong bag). Hình dưới: cách chiếu mặt đường vào ảnh.",
        "0.80 → 0.49 (PA2, camera trước).", figs=[PAINT, os.path.join(E, "x05_t3_ipm_390.png")], vertical=True)
problem("4.29", "Bộ lọc vạch sơn khoá sai khi xe đứng yên",
        "Thử trên Track01/02 với sai số biết trước: bản lọc đầu khoá sai khi xe dừng.",
        "Khi dừng, cùng một ảnh được đo hàng trăm lần – thực chất là một phép đo.",
        "Trọng số κ × min(1, quãng đường / 0.3 m) + sàn độ lệch chuẩn 0.3 m dọc / 0.15 m ngang.",
        "Track01/02 (chạy như không GNSS): 1.17 → 1.01 m, 1.32 → 1.16 m.", fig=os.path.join(E, "x06_t3_paint_steps_1100.png"))
s = new("4.30", "Vạch sơn: tham số nào quan trọng (mỗi file đổi một thứ)")
table(s, 0.55, TOP, 5.6, [["Thay đổi", "Public"], ["tay đòn ngang −0.25 / −0.10 / 0", "0.56 / 0.49 / 0.47"],
                           ["tay đòn dọc +0.15 / −0.15", "0.44 / 0.55"], ["κ 20 / 40 / 80 / 160", "0.44 / 0.39 / **0.37** / 0.38"],
                           ["yaw từ vạch sơn", "0.54 (kém yaw bộ bám)"], ["3 lượt đo", "0.50 / 0.38 (không lợi)"],
                           ["tìm kiếm trực tuyến", "lạc ở khúc cua"]], size=13, colw=[3.3, 2.3])
bullets(s, 0.55, TOP + 3.0, 5.6, 2.5, ["Giữ: κ 80, tay đòn dọc +0.15, ngang 0, yaw của bộ bám, 1 lượt (PE3)", "0.49 → **0.37** (hạng 2)",
                                         "TA cho chỉnh tham số Track03 theo điểm public"], size=14)
picture(s, os.path.join(E, "x06_t3_paint_steps_390.png"), 6.4, TOP, 6.4, 5.5)
problem("4.31", "Camera trước thấy ít mặt đường cạnh xe → thêm camera bên",
        "Camera trước chỉ nhìn 5–18 m phía trước.",
        "Tầm nhìn của một camera.",
        "Camera trước-trái / trước-phải (thông số ngoại từ yaml) cho phần mặt đường cạnh xe (< 5 m, |ngang| ≥ 1.5 m), khối NCC riêng, κ_side 40.",
        "0.37 → 0.33 (PF2); κ_side 80 → 0.47; 3 lượt với camera bên bị trôi (0.47).", fig=os.path.join(E, "x07_t3_side_cams.png"))
problem("4.32", "Vị trí dọc rung theo bộ bám",
        "Tốc độ suy ra từ kết quả nhảy 0–29 m/s trong khi xe chạy đều ~9.5 m/s; rung dọc 0.51 m ở frame private 409–950.",
        "Vạch làn chốt **ngang** ở mọi frame; đường NCC **dọc** phẳng giữa các vạch cắt ngang (hình trên) → vị trí dọc theo bộ bám.",
        "Giữa các vạch: pose đi theo tốc độ IMU × hệ số tỉ lệ tự học; chỉ khi ≥ 4 m/s và gần thẳng; κ_imu 10.",
        "0.33 → 0.28; rung dọc 0.51 → 0.24 m.",
        figs=[os.path.join(E, "x08_t3_ncc_profiles.png"), os.path.join(O, "t3_speed_jitter.png")], vertical=True)
problem("4.33", "Chuyển động IMU không có điều kiện lệch 1 m khi bò chậm / vào cua",
        "Bản không điều kiện lệch ~1 m ở frame 2400–2480; kéo về bộ bám ở mọi frame đứng yên làm lệch (NCC 0.54 → 0.28).",
        "Tốc độ IMU kém tin khi chậm và khi cua (điểm tham chiếu trượt); một lực kéo áp hàng trăm lần lúc đứng.",
        "Chỉ dùng khi v ≥ 4 m/s và |ω| ≤ 3°/s; lực kéo nhân theo quãng đường.",
        "đứng yên và yaw trùng từng byte bản trước; Q1–Q3 (κ_imu 20/10/5) đều 0.28.", fig=os.path.join(O, "t3_speed_jitter.png"))
problem("4.34", "Cuối Track03: vị trí trượt lùi 1.57 m khi phanh",
        "Frame 2850–3142 (~11 % private): PF2 / Q2 trượt từ +0.1 xuống −1.57 m so với bộ bám khi xe phanh 3 → 0 m/s.",
        "Đường NCC phẳng (không vạch chốt dọc), κ 80 cộng dồn một độ dốc nhỏ; camera trước và bên dốc ngược chiều; IMU × tỉ lệ đã học dự đoán −0.87 m.",
        "Giữ chế độ IMU khi giảm tốc tới 0.5 m/s nếu hệ số tỉ lệ đã học kỹ (sd ≤ 0.08) – tự tắt sau lần đứng 43 s.",
        "điểm dừng cuối −1.57 → −0.72 m; public 0.28 = 0.28 (Q4 = bài cuối).", fig=os.path.join(O, "t3_end_slide.png"))
problem("4.35", "Kaggle chỉ hiện 2 chữ số, cắt bớt",
        "Từ 21/10: 0.55496 hiện 0.55; κ_imu 5/10/20 và các bản Q7/Q8 đều 0.28.",
        "Ban tổ chức chống dò; độ phân giải ±0.005.",
        "Thước đo offline: độ khớp bản đồ (chỉ để loại cấu hình hỏng), tích phân IMU, chạy Track01/02 như không GNSS, camera bên làm phép thử độc lập (tương quan với public −0.77…−0.94).",
        "công cụ ra quyết định, không đổi điểm.", fig=os.path.join(E, "x08_t3_ncc_profiles.png"))
problem("4.36", "Còn mở: đầu Track03",
        "Camera trước không có ảnh ở frame 0–13; xe đang cua nhẹ ~4°/s nên chế độ IMU chỉ bật từ frame 21; bước của bộ bám frame 1–9 nhảy 5–17 m/s.",
        "Luồng h265 bắt đầu giữa một nhóm ảnh nén (khung khoá đầu ở message 14); ngưỡng 3°/s.",
        "Ứng viên (chưa nộp): --imu-wmax 5 (t3_R1) – frame 6–81 đi theo tốc độ IMU.",
        "đổi 0.037 m RMS ở public, lợi dự kiến ≤ 0.01.", fig=os.path.join(E, "x05_t3_ipm_390.png"))
L_table("4.37", "Những hướng đã thử và bỏ",
        [["Đã thử", "Kết quả"], ["Làm mượt GNSS ở Track02", "0.57111 → 0.57811 (xấu hơn)"], ["Yaw = hướng chạy ở Track01", "1.26235 → 1.35783"],
         ["Lệch thời gian ±0.1 s", "tối ưu −3.6 / −4.7 ms ≈ 0"], ["Hiệu chỉnh trượt trong bộ lọc yaw", "0.80359 → 0.80646"],
         ["Bộ ghép theo Hessian", "0.84 / 0.87 so với 0.80"], ["Kalman bền vững (Huber)", "đổi 1.9 cm"],
         ["Lệch giờ GNSS–IMU từ gia tốc kế", "không đo được (±75–220 ms)"], ["Bộ lọc hạt", "tốc độ sụp về 0 ở vùng phẳng"],
         ["Point cloud 3 camera trong bộ bám", "0.95 so với 0.91 m"], ["Yaw từ vạch sơn / hội tụ yaw nhanh", "0.54 so với 0.49; 0.85–1.08 so với 0.80"],
         ["Nhiều lượt vạch sơn có camera bên", "0.47 so với 0.33"]], size=12, colw=[6.5, 5.7])

# =================================================================== PHẦN 5
L_wide("5.1", "Track01/02: GNSS làm khung, chỉnh tham số theo điểm public",
       ["Chấm theo public toàn track (bạn đã xác nhận với TA 23/10); GNSS là output của một bộ định vị phẳng, riêng nó đã hơn baseline (1.31 / 0.57)",
        "Hằng số chọn theo điểm (được ghi rõ): τ yaw 7 / 15 s, trục gyro z, 2 lượt GICP, trừ lệch TB, tỉ lệ 0.8 / 1.646 m, σ yaw bản đồ 8 s (so với file dò), "
        "dời giờ GNSS −10 ms (T2)"], os.path.join(O, "t1t2_clean_chain.png"), text_h=1.45)
L_fig("5.2", "Track01/02: dùng bảng điểm như một phép đo (TA cho phép ở đây)",
      ["Điểm = P + 0.2·Y. Một cặp dời ±L: P₊² + P₋² − 2P² = 2L² → tách P và Y (hình: 3 lần nộp thật)",
       "Một thay đổi D bất kỳ: mean(e·D) = (P_D² − P² − mean|D|²)/2; tổ hợp tốt nhất a* = −G⁻¹m",
       "Dao động cosin 0.5 m·cos(πkt/T) trên x, y (bậc 0–18), yaw (bậc 1–8)",
       "**Dạy cho thuật toán:** yaw là lỗi lớn hơn; yaw đáp án là yaw thân xe; GNSS không lệch ngang → trừ lệch TB an toàn; yaw chậm 30 s → yaw bản đồ",
       "Kết quả: T1 0.46 → 0.36, T2 0.19 → 0.11 (hạng 1). Trọng số không có ý nghĩa vật lý, vô dụng với dữ liệu mới"],
      os.path.join(E, "x11_probe_parabola.png"), split=6.4)
L_text("5.3", "Track01/02 không nhân quả – và chúng em nói rõ",
       ["Pipeline Track01/02 dùng **RTS smoother** và làm mượt Gauss có tâm (dùng frame tương lai)",
        "Bài hạng 1 là **file dò bảng điểm**",
        "Văn bản luật 10/22: \"Offline trajectory corrections and leaderboard-based fitting are strictly prohibited\", \"Causal localization only\"",
        "Lời TA trên lớp 23/10 (theo bạn kể): Track01/02 chấm public → được dò; Track03 bắt buộc nhân quả",
        "Đã có pipeline nhân quả cho Track01/02 (run_causal_pipeline.sh) nhưng **chưa nộp**: lệch bản sạch 0.55 m / 0.82° (T1), 0.05 m / 0.71° (T2)",
        "[CẦN ĐIỀN: văn bản xác nhận của TA + file chọn \"Use for Final Score\"]"])
L_wide("5.4", "Track03: tối ưu phần mà điểm public không thấy",
       ["Public ≈ 13 % đầu (đường thẳng ~9.4 m/s); private gồm đèn đỏ 43 s (44 %), chạy lại + rẽ trái 88°, dừng cuối. Public rất nhạy với 84 frame đầu (giai đoạn MHT)",
        "Thay đổi cho private được đánh giá không qua bảng điểm: tích phân gia tốc kế, độ khớp bản đồ, tính nhất quán qua khúc cua, camera trước so với camera bên"],
       os.path.join(O, "t3_segments.png"), text_h=1.45)
L_text("5.5", "Track03: ổn định từ thiết kế",
       ["**Tất định:** chạy lại / đổi số luồng → trùng từng byte; chạy lại từ bộ đệm trống → trùng từng byte (Q4 kiểm chứng 25/10)",
        "**Vững với tham số:** sau nhiễu GICP bất đẳng hướng, mọi biến thể bám 96–99 % và khớp nhau ±0.5 m",
        "**Thử nơi có đáp án:** chạy Track01/02 như không có GNSS: lọc tốc độ giảm số lần lạc 49 → 9 / 80 → 33, sửa lúc chạy lại 32 → 15; tầng vạch sơn 1.17 → 1.01 m / 1.32 → 1.16 m",
        "**Cùng hiệu chuẩn cho mọi track** (cùng xe): pitch camera 1.0°, yaw −1.5°, cao 1.646 m",
        "Tầng mới chỉ đổi những frame nó đáng tin; phần còn lại trùng từng byte bản trước"])
L_fig("5.6", "Track03: chỉnh theo public, chọn theo private",
      ["Chỉnh theo public (TA cho phép): κ, tay đòn, κ_side, κ_imu – các hằng số tác dụng lên **mọi** frame, kể cả đèn đỏ",
       "Không nộp những thay đổi chỉ làm đổi giai đoạn khởi đầu",
       "**Bài cuối = t3_Q4:** public bằng Q2 (0.28) nhưng sửa trượt ở cuối (−1.57 → −0.72 m); Kaggle chỉ tính private cho bài được chọn → chọn tay",
       "Ước tính private (suy luận, không đo): **≈ 0.30–0.40** [CẦN ĐIỀN điểm thật sau 29/10]"],
      os.path.join(O, "t3_end_slide.png"), split=5.4)
L_text("5.7", "Track03: nhân quả, không hậu xử lý, chạy lại từ rosbag",
       ["Hàng k chỉ dùng: point cloud ≤ k, IMU ≤ thời điểm k, ảnh k, pose bộ bám tại k, trạng thái bộ lọc từ các frame < k",
        "Bản đồ cắt hành lang 30 m chỉ để đọc nhanh: frame k chỉ đọc bản đồ trong ~25 m quanh pose prior tại k",
        "Một script: bộ bám → giải mã h265 → bản đồ vạch sơn → đo → lọc; không file nào bị sửa tay",
        "**Đã kiểm chứng 25/10:** chạy lại với cấu hình Q4 → trùng từng byte với t3_Q4.csv. [CẦN ĐIỀN: đặt Q4 làm mặc định của script]"],
       code_text='SIDE=1 PASSES=1 \\\n FUSE="--kappa 80 --kappa-motion 0.3 --sd-floor 0.3,0.15 --lever-along 0.15 --lever-lat 0.0 --kappa-side 40" \\\n'
                 ' FUSE_EXTRA="--motion imu --imu-vmin 4 --imu-wmax 3 --kappa-imu 10 --imu-vstay 0.5 --imu-stay-sd 0.08" \\\n'
                 ' bash midterm01/tools/run_track3_paint.sh src/localization/result/track3/t3_Q4_rerun.csv', code_h=1.7)

# =================================================================== PHẦN 6
L_table("6.1", "Kết quả",
        [["", "Track01", "Track02", "Track03"], ["Baseline TA", "1.79075", "1.92587", "private ≤ 1.40"], ["Chỉ GNSS", "1.31094", "0.57111", "– (không GNSS)"],
         ["**Pipeline (thuật toán)**", "**0.51** clean_mapyaw_1", "**0.24** t2v_tsm10", "**0.28 public** t3_Q4 (nhân quả)"],
         ["Có dò bảng điểm", "**0.36** lb_t1_r5", "**0.11** lb_t2_r3", "không dùng"], ["Hạng (public 25/10)", "1 (đội 12: 0.40)", "1 (đội 12: 0.17)", "2 (đội 12: 0.24)"],
         ["Private", "–", "–", "[CẦN ĐIỀN sau 29/10]"], ["File chọn chấm điểm", "[CẦN ĐIỀN]", "[CẦN ĐIỀN]", "t3_Q4 (đề xuất)"]],
        size=14, colw=[2.8, 3.1, 3.1, 3.2])
s = new("6.2", "Quỹ đạo cuối trên bản đồ LiDAR")
picture(s, os.path.join(O, "trajectory_on_map_t3.png"), 0.45, TOP, 6.6, 5.55)
picture(s, os.path.join(O, "trajectory_on_map_t1.png"), 7.2, TOP, 5.6, 2.72)
picture(s, os.path.join(O, "trajectory_on_map_t2.png"), 7.2, TOP + 2.83, 5.6, 2.72)
L_table("6.3", "Ablation Track01/02: mỗi thành phần đáng giá bao nhiêu (Kaggle)",
        [["Thay đổi (mỗi lần một thứ)", "Track01", "Track02"], ["GNSS thay code mẫu", "602.39 → 1.31094", "– → 0.57111"],
         ["Kalman CTRV + RTS (+ lệch yaw b)", "1.26235 → 1.25338", "0.47281 → 0.40778"], ["GICP 2 lượt + trừ lệch (vị trí)", "0.905 → ~0.616 m", "xấu hơn (0.2267 → 0.414 m)"],
         ["Yaw gyro thay yaw tổ hợp", "1.559° → 0.865°", "0.330° → 0.249°"], ["Tầm GICP 40 → 20 m", "0.80 → 0.63", "–"],
         ["Bỏ GICP khi cua (+ tầm 20)", "0.63 → 0.62", "–"], ["CSM bản đồ", "0.62 → 0.55", "0.28 → 0.25"],
         ["Yaw bản đồ (σ 8 s)", "0.55 → 0.51", "xấu hơn (offline)"], ["Dời giờ GNSS −10 ms", "–", "0.25 → 0.24"]],
        size=12, colw=[3.4, 1.9, 1.9], fig=os.path.join(O, "t1t2_clean_chain.png"), fig_w=4.7)
L_table("6.4", "Ablation Track03: các tầng và tham số (điểm public)",
        [["Bỏ / đổi", "Public"], ["Bài cuối t3_Q4", "**0.28**"], ["bỏ sửa đoạn cuối (Q2)", "0.28 (chỉ khác private)"], ["bỏ chuyển động IMU (PF2)", "0.33"],
         ["bỏ camera bên (PE3)", "0.37"], ["κ 80 → 40 / 20", "0.39 / 0.44"], ["bỏ toàn bộ vạch sơn (Z3a)", "0.80"],
         ["bỏ hướng trọng lực cố định (W3h)", "0.89"], ["bỏ hiệu chuẩn lại tay đòn (L2)", "1.05"], ["bộ bám trước CSM / sửa gyro (A2h)", "0.96"]],
        items=["Offline: CSM 69 → 98 % chấp nhận; sửa lúc chạy lại 2.46 → 0.72 m; rung dọc 0.51 → 0.24 m"],
        size=12, colw=[4.1, 2.6], fig=os.path.join(O, "t3_layers.png"), fig_w=5.2)
L_text("6.5", "Phân công",
       ["Đội 1 thành viên → 100 % [CẦN ĐIỀN: họ tên]",
        "Sinh viên: chọn hướng đi, duyệt mọi lần nộp Kaggle, trao đổi với TA (20/10, 23/10), tải dữ liệu, cài môi trường, thuyết trình [CẦN ĐIỀN: phần tự làm]",
        "Trợ lý AI (Claude Code): viết công cụ, chạy thí nghiệm, phân tích và ghi nhật ký theo chỉ dẫn; không bao giờ tự nộp",
        "[CẦN ĐIỀN: mô tả phần dùng AI theo yêu cầu môn học]"])
L_text("6.6", "Bài học",
       ["Kiểm tra đơn vị và hệ trục trước: mất ~70 phút chạy (baseline, NDT ×2–×4) trước khi kiểm tra tỉ lệ",
        "Kiểm tra thước đo có thiên vị không (khoảng cách trung bình thiên về tỉ lệ nhỏ)",
        "Đầu vào giới hạn kết quả: thuật toán khớp sai 5–8 mm trên point cloud hoàn hảo → sửa đầu vào, thêm thông tin mới (vạch sơn), không cần bộ lọc phức tạp hơn",
        "Giả thuyết đúng về vật lý vẫn có thể sai với đáp án (bù trễ yaw GNSS) – phải thử",
        "Nhiều lần nhìn cùng một khung cảnh = một phép đo (đứng yên) – đúng cho GICP, CSM và vạch sơn",
        "Tự đọc lại code để kiểm tra nhân quả: lỗi ghi ngược MHT trông như điểm tốt",
        "Ghi chép mọi thứ, mọi file tạo lại được (34/34 trùng từng byte)"])
L_text("6.7", "Hướng tiếp theo",
       ["Track03: dùng cấu trúc 3D của bản đồ (cột, mép lề, tường) chiếu vào ảnh để chốt vị trí dọc ở chỗ không có vạch sơn (thử ngày 26/10) [CẦN ĐIỀN: kết quả]",
        "Đầu Track03: frame 0–13 không có ảnh camera trước, bật IMU từ frame 6 (t3_R1, chưa nộp)",
        "Track01/02: đánh giá bản nhân quả (run_causal_pipeline.sh) trên Kaggle",
        "Mô hình sai độ sâu theo khoảng cách cho LoGeR (p' = p(1 + 4c)/(1 + c·r)) thay cho một tỉ lệ"])
dark("6.8", "Cảm ơn – câu hỏi",
     ["Track03 có dùng frame tương lai không? – Không (slide 4.20, 5.7)", "Vì sao được dò ở Track01/02? – TA cho phép trên lớp 23/10 (slide 5.3)",
      "Vì sao không dùng NDT? – kém GNSS; GICP nhanh hơn ~200× (slide 2.5)", "Điểm private Track03? – chưa biết; ước tính 0.3–0.4 (slide 5.6)"])

out = os.path.join(D, "presentation_full_vi.pptx")
prs.save(out)
print("wrote", out, len(prs.slides), "slides")

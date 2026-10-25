"""Midterm01 presentation (4-6 min, 29/10) -> docs/report/presentation_midterm01.pptx  (python-pptx, native charts).

8 slides, English, speaker notes in the notes pane (Vietnamese translation: docs/report/loi_noi_thuyet_trinh.md).
Figures: fig/paint_matching.png (make_paint_figure.py), fig/t3_speed.csv (Track03 speed series of t3_PF2 / t3_Q4).
"""
import csv
import os

from pptx import Presentation
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

D = os.path.dirname(os.path.abspath(__file__))
NAVY = RGBColor(0x14, 0x24, 0x2E)
TEAL = RGBColor(0x13, 0x7C, 0x8B)
TEAL_L = RGBColor(0xD8, 0xEE, 0xF0)
AMBER = RGBColor(0xE8, 0x9B, 0x2B)
INK = RGBColor(0x1E, 0x2A, 0x32)
MUTED = RGBColor(0x5F, 0x6F, 0x7A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CARD = RGBColor(0xF3, 0xF6, 0xF8)
ICE = RGBColor(0xB9, 0xD9, 0xDE)
# chart colours (validated pair: CVD separation 16.6, contrast >= 3:1 on white) + a neutral for reference marks
CH_TEAL = RGBColor(0x00, 0x8A, 0x9E)
CH_AMBER = RGBColor(0xC9, 0x78, 0x1A)
CH_SLATE = RGBColor(0x8A, 0x99, 0xA3)
HEAD, BODY = "Cambria", "Calibri"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


def bg(slide, color):
    f = slide.background.fill
    f.solid()
    f.fore_color.rgb = color


def text(slide, x, y, w, h, runs, size=16, color=INK, bold=False, font=BODY, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, 0)
    paras = runs if isinstance(runs, list) else [runs]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(6)
        for seg in (para if isinstance(para, list) else [para]):
            t, o = (seg, {}) if isinstance(seg, str) else seg
            r = p.add_run()
            r.text = t
            r.font.name = o.get("font", font)
            r.font.size = Pt(o.get("size", size))
            r.font.bold = o.get("bold", bold)
            r.font.color.rgb = o.get("color", color)
    return tb


def title(slide, t, sub=None):
    text(slide, 0.6, 0.45, 12.1, 0.8, t, size=32, bold=True, font=HEAD, color=NAVY)
    if sub:
        text(slide, 0.6, 1.17, 12.1, 0.5, sub, size=16, color=MUTED)


def card(slide, x, y, w, h, fill=CARD):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.adjustments[0] = 0.08
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.fill.background()
    s.shadow.inherit = False
    return s


def badge(slide, x, y, label, fill=TEAL, d=0.62, size=18):
    s = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.fill.background()
    tf = s.text_frame
    tf.margin_left = tf.margin_right = 0
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    r.font.size, r.font.bold, r.font.name = Pt(size), True, HEAD
    r.font.color.rgb = WHITE
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE


def notes(slide, t):
    slide.notes_slide.notes_text_frame.text = t


def footer(slide, n):
    text(slide, 0.6, 7.0, 8, 0.3, "SDC 2026 Midterm01 · Team 11", size=10, color=MUTED)
    text(slide, 11.9, 7.0, 0.8, 0.3, str(n), size=10, color=MUTED, align=PP_ALIGN.RIGHT)


def chart_title(ch, ttl):
    ch.has_title = True
    ch.chart_title.text_frame.text = ttl
    r0 = ch.chart_title.text_frame.paragraphs[0].runs[0]
    r0.font.size, r0.font.bold, r0.font.name = Pt(14), True, BODY
    r0.font.color.rgb = INK


def bar_chart(slide, x, y, w, h, cats, vals, ttl, colors):
    """one series, data labels on the bars, no value axis; colors: one per bar"""
    cd = CategoryChartData()
    cd.categories = cats
    cd.add_series("score", vals)
    ch = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(x), Inches(y), Inches(w), Inches(h), cd).chart
    chart_title(ch, ttl)
    ch.has_legend = False
    pl = ch.plots[0]
    pl.gap_width = 60
    pl.has_data_labels = True
    dl = pl.data_labels
    dl.number_format, dl.number_format_is_linked = "0.00", False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.font.size, dl.font.name = Pt(12), BODY
    dl.font.color.rgb = INK
    ser = pl.series[0]
    for i, col in enumerate(colors):
        pt = ser.points[i]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = col
    va = ch.value_axis
    va.has_major_gridlines = False
    va.visible = False
    va.minimum_scale = 0
    ca = ch.category_axis
    ca.tick_labels.font.size, ca.tick_labels.font.name = Pt(10), BODY
    ca.tick_labels.font.color.rgb = MUTED
    ca.format.line.color.rgb = RGBColor(0xC8, 0xD0, 0xD6)
    return ch


n = 0
# 1 ---------------------------------------------------------------- title
s = prs.slides.add_slide(BLANK)
bg(s, NAVY)
text(s, 0.8, 2.0, 11.5, 1.4, "Localizing a car against a LiDAR map", size=44, bold=True, font=HEAD, color=WHITE)
text(s, 0.8, 3.35, 11.5, 0.8, "Camera point clouds, road paint and the IMU — Tracks 01, 02 and the GNSS-free Track 03",
     size=20, color=ICE)
text(s, 0.8, 5.6, 11.5, 0.5, "SDC 2026 · Midterm01 · Team 11 (solo, AI-assisted)", size=16, color=AMBER, bold=True)
notes(s, "Hello, I'm team 11 and I worked alone, with an AI coding assistant. The task is to estimate the car's x, y and yaw "
         "for every camera frame against a pre-built LiDAR map. I'll show the pipeline for Tracks 01 and 02 and how I used "
         "the leaderboard there; most of the time goes to Track 03, which has no GNSS.")

# 2 ---------------------------------------------------------------- task
n += 1
s = prs.slides.add_slide(BLANK)
bg(s, WHITE)
title(s, "The task and the score", "Score = position RMSE (m) + 0.2 × yaw RMSE (deg) — lower is better")
items = [("01", "Track 01", "3011 frames · GNSS + IMU"),
         ("02", "Track 02", "3299 frames · GNSS + IMU · 13 s standstill at the start"),
         ("03", "Track 03", "3142 frames · no GNSS · causal only · graded on the hidden private part (~87 %)")]
for i, (b, h, d_) in enumerate(items):
    y = 2.0 + i * 1.45
    card(s, 0.6, y, 7.4, 1.2)
    badge(s, 0.85, y + 0.29, b)
    text(s, 1.75, y + 0.18, 6.0, 0.4, h, size=20, bold=True, font=HEAD, color=NAVY)
    text(s, 1.75, y + 0.6, 6.0, 0.5, d_, size=14, color=MUTED)
card(s, 8.5, 2.0, 4.2, 4.1, TEAL_L)
text(s, 8.85, 2.25, 3.6, 0.5, "Inputs", size=20, bold=True, font=HEAD, color=TEAL)
text(s, 8.85, 2.85, 3.6, 3.1, ["Point cloud per frame from images (LoGeR, unknown scale)",
                                 "LiDAR map with intensity, 57 M points", "3 cameras (h265 video in the bag)",
                                 "IMU 100 Hz, noisy initial pose", "GNSS (Tracks 01/02 only)"], size=15, color=INK)
footer(s, n)
notes(s, "There are three tracks. Tracks 01 and 02 have GNSS and IMU. Track 03 has no GNSS, the localization must be causal, "
         "and it is graded on a hidden private part, about 87 percent of the frames. The score adds position RMSE in metres "
         "and 0.2 times yaw RMSE in degrees, so one degree of yaw costs as much as 20 centimetres.")

# 3 ---------------------------------------------------------------- T1/T2 pipeline + chart
n += 1
s = prs.slides.add_slide(BLANK)
bg(s, WHITE)
title(s, "Tracks 01/02: GNSS gives the frame, the map refines it")
steps = [("0", "Scale ×9, not ×3", "camera height above the road = 1.646 m"),
         ("1", "GNSS + Kalman (CTRV) + RTS", "smooth, but GNSS heading lags 1.1–1.4 s"),
         ("2", "GICP vs map, range ≤ 20 m", "image depth degrades with range"),
         ("3", "Correlative map matching", "fixes the along-road position"),
         ("4", "Yaw: gyro + map yaw", "gyro for fast, map for slow changes")]
for i, (b, h, d_) in enumerate(steps):
    y = 1.45 + i * 1.07
    card(s, 0.6, y, 5.4, 0.9)
    badge(s, 0.78, y + 0.17, b, fill=AMBER if b == "0" else TEAL, d=0.56, size=16)
    text(s, 1.55, y + 0.1, 4.3, 0.4, h, size=16, bold=True, color=NAVY)
    text(s, 1.55, y + 0.47, 4.3, 0.4, d_, size=13, color=MUTED)
bar_chart(s, 6.4, 1.45, 6.4, 5.3, ["Baseline", "GNSS +\nKalman", "+ GICP", "+ range\n20 m", "+ map\nmatching", "+ map\nyaw"],
          [1.79, 1.26, 0.79, 0.63, 0.55, 0.51], "Track 01 pipeline score (Kaggle)", [CH_SLATE] + [CH_TEAL] * 5)
footer(s, n)
notes(s, "For Tracks 01 and 02 the first finding was the scale: the baseline scaled the point clouds by 3, but the camera "
         "height above the road says about 9, so every frame is rescaled by the road height. Then GNSS goes through a Kalman "
         "filter and RTS smoother; GICP against the map refines the position, only within 20 metres because image depth gets "
         "worse with range; correlative map matching fixes the along-road position; and the yaw comes from the gyro, with the "
         "slow part from the map. Track 01 went from 1.79 to 0.51, Track 02 from 1.93 to 0.24.")

# 4 ---------------------------------------------------------------- leaderboard as a measurement
n += 1
s = prs.slides.add_slide(BLANK)
bg(s, WHITE)
title(s, "Tracks 01/02: the leaderboard as a measurement", "Allowed by the TA for Tracks 01/02 only — never used on Track 03")
card(s, 0.6, 1.85, 6.3, 2.05, TEAL_L)
text(s, 0.9, 2.05, 5.8, 0.5, "The score is a quadratic oracle", size=20, bold=True, font=HEAD, color=TEAL)
text(s, 0.9, 2.65, 5.8, 1.2, ["Add a known change D to a file, read its score P_D:",
                               "mean(e·D) = (P_D² − P² − mean|D|²) / 2"], size=16, color=INK)
text(s, 0.6, 4.2, 6.3, 2.6, [
    [("What it taught us: ", {"bold": True, "color": NAVY}), "yaw is the larger error, and the pipeline yaw had a slow "
     "error with a ~30 s period → fixed in the pipeline with map yaw (0.55 → 0.51)."],
    [("Kept apart: ", {"bold": True, "color": NAVY}), "the probed files are labelled as such; the pipeline numbers are our "
     "algorithm's result."]], size=15)
cd = CategoryChartData()
cd.categories = ["Track 01", "Track 02"]
cd.add_series("Pipeline (algorithm)", (0.51, 0.24))
cd.add_series("With leaderboard probing", (0.36, 0.11))
ch = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(7.3), Inches(1.85), Inches(5.5), Inches(4.9), cd).chart
chart_title(ch, "Score — both rank 1 with probing")
ch.has_legend = True
ch.legend.position = XL_LEGEND_POSITION.BOTTOM
ch.legend.include_in_layout = False
ch.legend.font.size, ch.legend.font.name = Pt(11), BODY
pl = ch.plots[0]
pl.gap_width = 70
pl.overlap = -8
pl.has_data_labels = True
pl.data_labels.number_format, pl.data_labels.number_format_is_linked = "0.00", False
pl.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
pl.data_labels.font.size, pl.data_labels.font.name = Pt(12), BODY
for ser, col in zip(pl.series, (CH_TEAL, CH_AMBER)):
    ser.format.fill.solid()
    ser.format.fill.fore_color.rgb = col
ch.value_axis.visible = False
ch.value_axis.has_major_gridlines = False
ch.value_axis.minimum_scale = 0
ch.category_axis.tick_labels.font.size = Pt(12)
ch.category_axis.tick_labels.font.color.rgb = MUTED
footer(s, n)
notes(s, "The TA allowed using the leaderboard on Tracks 01 and 02. Because the score is an RMSE, adding a known change to a "
         "file and reading the new score measures how the hidden error projects onto that change. That showed that yaw was "
         "the larger error and that the pipeline yaw had a slow 30-second error, which I then fixed in the pipeline with map "
         "yaw. With probing, both tracks rank first, 0.36 and 0.11; the pipeline alone gives 0.51 and 0.24, and the report "
         "keeps the two apart; the one pipeline constant chosen against a probed file, the map-yaw time constant, is "
         "labelled too. I never did this on Track 03.")

# 5 ---------------------------------------------------------------- T3 overview: three layers
n += 1
s = prs.slides.add_slide(BLANK)
bg(s, WHITE)
title(s, "Track 03: no GNSS, causal — three layers", "Every output row uses only the current and past frames")
layers = [("1", CH_SLATE, "LiDAR-map tracker", "yaw hypotheses at the start · IMU speed filter · gyro on gravity (IMU "
           "tilted 6°) · scan matching on the map ±15 m", "0.80"),
          ("2", CH_TEAL, "Road-paint matching", "3 cameras vs the paint in the map's LiDAR intensity", "0.33"),
          ("3", CH_AMBER, "IMU along-motion", "smooth speed × learnt scale between road markings",
           "0.28")]
for i, (b, col, h, d_, sc) in enumerate(layers):
    y = 1.85 + i * 1.6
    card(s, 0.6, y, 6.1, 1.4)
    badge(s, 0.82, y + 0.39, b, fill=col)
    text(s, 1.7, y + 0.15, 3.7, 0.45, h, size=18, bold=True, font=HEAD, color=NAVY)
    text(s, 1.7, y + 0.6, 3.75, 0.8, d_, size=13, color=MUTED)
    text(s, 5.45, y + 0.3, 1.1, 0.8, sc, size=28, bold=True, font=HEAD, color=NAVY, align=PP_ALIGN.RIGHT,
         anchor=MSO_ANCHOR.MIDDLE)
text(s, 0.6, 6.62, 6.1, 0.3, "Public score after each layer · constants tuned on it (allowed) · outputs never edited",
     size=11, color=MUTED)
bar_chart(s, 7.0, 1.7, 5.8, 5.1, ["First\ncausal", "Tracker\nfixes", "Paint,\nfront cam", "Tuned κ,\nlever", "+ side\ncams",
                                  "+ IMU\nmotion"],
          [0.96, 0.80, 0.49, 0.37, 0.33, 0.28], "Track 03 public score (Kaggle)",
          [CH_SLATE, CH_SLATE, CH_TEAL, CH_TEAL, CH_TEAL, CH_AMBER])
footer(s, n)
notes(s, "Track 03 is built in three layers, and every output uses only the current and past frames. First, a LiDAR-map "
         "tracker: several yaw hypotheses at the start, an IMU speed filter, the gyro projected on gravity because the IMU "
         "is tilted 6 degrees, and scan matching on the map. Debugging it with physics instead of ground truth brought the "
         "public score to 0.80. The two layers on top use the cameras and took it to 0.28. Their constants were tuned on the "
         "public score, which the TA allows; the output files are never edited.")

# 6 ---------------------------------------------------------------- road paint
n += 1
s = prs.slides.add_slide(BLANK)
bg(s, WHITE)
title(s, "Layer 2: the cameras see the paint that the LiDAR map stores")
steps = [("Decode", "front + front-left / front-right h265 video"),
         ("Project", "the map's road surface into each fisheye image (IPM)"),
         ("Extract", "thin bright structures = paint, in camera and map"),
         ("Correlate", "NCC over ±2 m × ±1.5 m shifts (FFT)"),
         ("Fuse", "causal error-state filter on top of the tracker")]
for i, (h, d_) in enumerate(steps):
    y = 1.5 + i * 1.02
    card(s, 0.6, y, 5.2, 0.86)
    badge(s, 0.78, y + 0.16, str(i + 1), d=0.54, size=15)
    text(s, 1.5, y + 0.09, 4.2, 0.4, h, size=16, bold=True, color=NAVY)
    text(s, 1.5, y + 0.45, 4.2, 0.4, d_, size=13, color=MUTED)
text(s, 0.6, 6.68, 5.2, 0.3, [[("Public 0.80 → 0.33", {"bold": True, "color": TEAL}), "   (red light, frame 1100 →)"]],
     size=13)
s.shapes.add_picture(os.path.join(D, "fig", "paint_matching.png"), Inches(6.15), Inches(1.3), height=Inches(5.55))
footer(s, n)
notes(s, "The LiDAR map also stores intensity, and road paint is bright in it: lane lines, stop lines, crosswalks, accurate to "
         "a few centimetres. The bag contains three camera videos. For every frame I project the map's road surface into "
         "each fisheye image, keep only thin bright structures, and correlate this camera paint with the map paint over shifts "
         "of plus or minus 2 metres. That gives a likelihood for the car's position, which a causal error-state filter "
         "combines with the tracker. Here, at a red light, the tracker was 0.66 metres off; after matching, the waiting box "
         "lines up. Public score 0.80 to 0.33.")

# 7 ---------------------------------------------------------------- IMU motion
n += 1
s = prs.slides.add_slide(BLANK)
bg(s, WHITE)
title(s, "Layer 3: the IMU carries the along position")
text(s, 0.6, 1.5, 5.4, 4.6, [
    [("Problem: ", {"bold": True, "color": NAVY}), "lane lines pin the lateral position every frame; the along position "
     "only at dashes, stop lines and crosswalks. In between it followed the tracker's jitter."],
    [("Fix: ", {"bold": True, "color": NAVY}), "between markings the pose moves with the IMU speed × a scale the filter "
     "learns (odometry is 10–15 % off); only at ≥ 4 m/s and nearly straight, kept on while braking to a stop."]],
     size=16)
for i, (big, lab) in enumerate([("0.51 → 0.24 m", "along jitter, private frames 409–950"), ("0.33 → 0.28", "public score")]):
    y = 4.35 + i * 1.25
    card(s, 0.6, y, 5.4, 1.05)
    text(s, 0.9, y + 0.12, 2.9, 0.8, big, size=26, bold=True, font=HEAD, color=CH_TEAL, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 3.85, y + 0.12, 2.0, 0.8, lab, size=14, color=MUTED, anchor=MSO_ANCHOR.MIDDLE)
rows = list(csv.reader(l for l in open(os.path.join(D, "fig", "t3_speed.csv")) if not l.startswith("#")))[1:]
xy = XyChartData()
for name, col in (("Paint only", 2), ("Paint + IMU motion", 3), ("IMU speed × learnt scale", 4)):
    se = xy.add_series(name)
    for r in rows:
        se.add_data_point(float(r[0]), float(r[col]))
ch = s.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER_LINES_NO_MARKERS, Inches(6.3), Inches(1.4), Inches(6.5), Inches(5.35),
                        xy).chart
chart_title(ch, "Speed implied by the output, 10 s of Track 03 (m/s)")
ch.has_legend = True
ch.legend.position = XL_LEGEND_POSITION.BOTTOM
ch.legend.include_in_layout = False
ch.legend.font.size, ch.legend.font.name = Pt(11), BODY
for ser, col, wpt, dash in zip(ch.plots[0].series, (CH_AMBER, CH_TEAL, CH_SLATE), (1.0, 2.25, 1.5),
                               (None, None, MSO_LINE.DASH)):
    ser.smooth = False
    ln = ser.format.line
    ln.color.rgb = col
    ln.width = Pt(wpt)
    if dash:
        ln.dash_style = dash
va = ch.value_axis
va.minimum_scale, va.maximum_scale, va.major_unit = 0, 30, 10
va.has_major_gridlines = True
va.major_gridlines.format.line.color.rgb = RGBColor(0xE3, 0xE8, 0xEB)
va.format.line.fill.background()
va.tick_labels.font.size, va.tick_labels.font.color.rgb = Pt(11), MUTED
ca = ch.category_axis
ca.minimum_scale, ca.maximum_scale, ca.major_unit = 0, 10, 2
ca.has_major_gridlines = False
ca.tick_labels.font.size, ca.tick_labels.font.color.rgb = Pt(11), MUTED
ca.format.line.color.rgb = RGBColor(0xC8, 0xD0, 0xD6)
ca.has_title = True
ca.axis_title.text_frame.text = "time (s)"
at = ca.axis_title.text_frame.paragraphs[0].runs[0]
at.font.size, at.font.bold, at.font.color.rgb = Pt(11), False, MUTED
footer(s, n)
notes(s, "One problem remained. Lane lines fix the lateral position in every frame, but the along-road position is only "
         "fixed at dashes, stop lines and crosswalks. In between, the estimate followed the tracker's jitter: the implied "
         "speed jumped between 0 and 29 metres per second while the car drove a steady 9.5. So between markings the pose now "
         "moves with the IMU speed times a scale that the filter learns, because the odometry is 10 to 15 percent off. It is "
         "used only when the car is fast and nearly straight, and it stays on while braking to a stop. The speed is now "
         "smooth, and the public score went from 0.33 to 0.28.")

# 8 ---------------------------------------------------------------- results + lessons (dark)
n += 1
s = prs.slides.add_slide(BLANK)
bg(s, NAVY)
text(s, 0.6, 0.45, 12.1, 0.8, "Results and what we would keep", size=32, bold=True, font=HEAD, color=WHITE)
tiles = [("0.36", "Track 01 · rank 1", "pipeline alone 0.51"), ("0.11", "Track 02 · rank 1", "pipeline alone 0.24"),
         ("0.28", "Track 03 · rank 2 (public)", "causal pipeline, no GNSS")]
for i, (big, lab, sub) in enumerate(tiles):
    x = 0.6 + i * 4.15
    card(s, x, 1.5, 3.85, 2.15, RGBColor(0x22, 0x3A, 0x47))
    text(s, x + 0.3, 1.65, 3.3, 1.0, big, size=50, bold=True, font=HEAD, color=AMBER)
    text(s, x + 0.3, 2.75, 3.3, 0.4, lab, size=16, bold=True, color=WHITE)
    text(s, x + 0.3, 3.15, 3.3, 0.4, sub, size=14, color=ICE)
text(s, 0.6, 4.1, 12.1, 2.8, [
    [("Keep: ", {"bold": True, "color": AMBER}), "physics checks instead of guessing — the right scale, short-range "
     "matching, the IMU mounting, standstills, and the road paint as a centimetre-level landmark."],
    [("Next: ", {"bold": True, "color": AMBER}), "the map's 3D structure (poles, curbs, walls) in the camera images, for the "
     "along position where the road has no paint."],
    [("Honesty: ", {"bold": True, "color": AMBER}), "probed files labelled; Track 03 is pipeline only and reruns from raw data "
     "with one script."]], size=18, color=WHITE)
notes(s, "Results: Track 01 0.36 and Track 02 0.11, both first, with the pipeline alone at 0.51 and 0.24. Track 03 is second "
         "on the public part with 0.28, from a purely causal pipeline. What worked was physics: the right scale, short-range "
         "matching, the IMU mounting, and road paint as a landmark. Next I would use the map's 3D structure, like poles and "
         "curbs, for the along position. Thank you.")

out = os.path.join(D, "presentation_midterm01.pptx")
prs.save(out)
print("wrote", out)
